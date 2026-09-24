"""Create a non-destructive group-meeting update from the corrected v6 deck."""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "cow_rr_group_meeting_20260924_v6_corrected.pptx"
TARGET = ROOT / "cow_rr_group_meeting_20260924_v7_method_verified.pptx"


def replace_shape(slide, old, new):
    found = [shape for shape in slide.shapes if shape.has_text_frame and shape.text == old]
    if len(found) != 1:
        raise ValueError(f"Expected one text shape, got {len(found)}: {old!r}")
    paragraphs = found[0].text_frame.paragraphs
    runs = [run for paragraph in paragraphs for run in paragraph.runs]
    if not runs:
        raise ValueError(f"Shape has no text run: {old!r}")
    runs[0].text = new
    for run in runs[1:]:
        run.text = ""


def main():
    if TARGET.exists():
        raise FileExistsError(TARGET)
    deck = Presentation(SOURCE)
    if len(deck.slides) != 18:
        raise ValueError("Unexpected source deck size")
    edits = {
        1: {
            "组会路线：整曲线质量选择＋约束峰判；残差校正列后续候选":
                "组会路线已做同输入核证：主方法优势未证实；残差校正仍列后续候选",
        },
        4: {
            "B73-I：73段6–20秒；久福：271个30秒窗（216事件参考、173 RR配对）；林甸49段仅同场敏感性分析。":
                "林甸R2：47个约30秒完整窗做同场主配对核证，2窗参考未完成；73短片和久福分账。",
            "RR=60×N/T。73段按各自窗口时长；久福按冻结的30秒标注副本计时。拒测和不可观察均单列覆盖，不计作零次。":
                "RR=60×N/T。林甸47窗用R2冻结时长；73短片逐段计时；久福用30秒副本。拒测单列。",
        },
        13: {
            "47窗同口径对照：增益尚未证实": "47窗原文信号规则复放：增益未证实",
            "47个林甸约30秒窗，统一R2参考；r20是项目基线，不等于忠实Chen复现":
                "47个林甸约30秒R2完整窗；同一r20温度输入与冻结时长",
            "r20：R² .8728；MAE 3.020": "C0P0：R² .7592；MAE 3.362",
            "F1G1P1：R² .8909；MAE 2.978": "C1P0：R² .8499；MAE 2.935",
            "F1G0P1：R² .8943；MAE 2.935": "C1P1：R² .8662；MAE 2.935",
            "两种配置的配对ΔMAE区间均跨0": "ΔMAE −.428；95%CI跨0",
            "不能声称稳定优于原文方法": "exact 15→13；F1 .758→.762",
            "先完成忠实原文复现对照。": "预定判据未通过；不主张稳定方法增益。",
            "47窗与73短片不合并；E040的73段消融引用了调整后的计数，不能作干净主证据。":
                "C1P1−C0P0 配对ΔMAE=−.428，95%CI[−1.834,.722]；仅信号规则复放，非Chen端到端。",
        },
        15: {
            "差异在整曲线质量选择与约束峰值判定；方法优势尚待忠实复现对照":
                "同输入信号规则复放后，方法优势未证实；原作者端到端权重仍缺",
        },
        16: {
            "保留原组会路线；论文创新主张仍需同输入、同参考对照":
                "47窗同输入R2核证已完成；方法增益未通过预定判据",
            "47窗同口径F1G1P1对r20：ΔMAE区间跨0。":
                "47窗C1P1对C0P0：ΔMAE区间跨0，exact 15→13。",
            "73段逐段6–20秒；久福30秒标注副本；49段仅作内部敏感性分析":
                "47窗R2主配对；73段B73-I次分析；久福仅事后迁移。",
        },
        18: {
            "先核实Chen忠实复现，再做同输入/参考配对对照。":
                "同输入47窗主比较CI跨0且完全计数下降；方法增益未证实。",
            "73段B73-I：已实现基线R²=0.928992（6–20秒）。":
                "73段旧流程B73-I R²=.929；新四组仅作次分析，不混为一法。",
        },
    }
    for page, replacements in edits.items():
        for old, new in replacements.items():
            replace_shape(deck.slides[page - 1], old, new)
    chart = deck.slides[12].shapes[3].chart
    if [str(category.label) for category in chart.plots[0].categories] != ["项目r20基线", "F1G1P1", "F1G0P1"]:
        raise ValueError("Unexpected source chart categories")
    values = CategoryChartData()
    values.categories = ["C0P0", "C1P0", "C1P1"]
    values.add_series("RR R² × 100", [75.9195, 84.9910, 86.6151])
    chart.replace_data(values)
    deck.save(TARGET)
    check = Presentation(TARGET)
    assert len(check.slides) == 18
    for page, replacements in edits.items():
        content = "\n".join(s.text for s in check.slides[page - 1].shapes if s.has_text_frame)
        for new in replacements.values():
            if new not in content:
                raise AssertionError(f"Missing new text on slide {page}: {new}")
    assert [str(c.label) for c in check.slides[12].shapes[3].chart.plots[0].categories] == ["C0P0", "C1P0", "C1P1"]
    print(TARGET)


if __name__ == "__main__":
    main()
