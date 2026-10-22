"""Import independently rebuilt metric snapshots; preserve completed experiments."""
from __future__ import annotations

import importlib.util
import json
import math
import os
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
OLD = ROOT / "Experiment/07_all_video_pose_retrain"
NEW_REPLAY = HERE / "rr_replay"
OLD_REPLAY = OLD / "replay_lindian_20261016_v1"
NEW_TRAINING = HERE / "training_runs/yolo11n_pose_frame_random_20261019_v1"
OLD_TRAINING = OLD / "training_runs/yolo11n_pose_lindian_20261016_v1"
OUT = HERE / "metric_tables_20261021_v2"
os.environ["YOLO_CONFIG_DIR"] = str(HERE / "runtime/metric_export")
sys.dont_write_bytecode = True
sys.path.insert(0, str(OLD))

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from export_metric_tables import add_sheet, excel_value, markdown_table, same, sha256

ARMS = ("C0P0", "C1P0", "C0P1", "C1P1")
LABELS = {"C0P0": "原文融合+原文峰规则", "C1P0": "质量选择+原文峰规则",
          "C0P1": "原文融合+约束峰规则", "C1P1": "质量选择+约束峰规则"}
SCOPES = ("short73", "anchored49", "P1", "P2", "P3", "all82", "all155")
SCOPE_LABELS = {"short73": "旧73原时长参考", "anchored49": "49窗中的完整参考", "P1": "P1", "P2": "P2", "P3": "P3",
                "all82": "仅约30秒合计82", "all155": "全部合计155（混合时长）"}
METRICS = (("rr_mae_bpm", "RR MAE", "次/分，越低越好"),
           ("rr_rmse_bpm", "RR RMSE", "次/分，越低越好"), ("rr_r2", "RR R²", "1-SSE/SST，越高越好"),
           ("mean_count_accuracy", "平均相对计数准确度", "逐窗平均，非事件准确率"),
           ("exact_display", "完全计数相同/可评分", "窗/窗"), ("exact_count_rate", "完全计数率", "完全计数相同窗比例"),
           ("tp", "事件TP", "次"), ("fp", "事件FP", "次"), ("fn", "事件FN", "次"),
           ("event_f1", "事件F1", "仅完整事件参考，±0.30秒"))


def save(name, frame):
    with (OUT / name).open("x", encoding="utf-8-sig", newline="") as stream:
        pd.DataFrame(frame).to_csv(stream, index=False)


def json_file(name, value):
    with (OUT / name).open("x", encoding="utf-8") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def select(frame, scope, historical="all"):
    if scope == "all82":
        selected = frame[frame.cohort.ne("short73")]
    elif scope == "all155":
        selected = frame
    else:
        selected = frame[frame.cohort.eq(scope)]
    return selected if historical == "all" else selected[selected.historical_group.eq(historical)]


def measure(frame):
    scored = frame[frame.eligible_rr]
    events = scored[scored.eligible_events]
    n, e = len(scored), len(events)
    tp, fp, fn = (float(events[key].sum()) if e else np.nan for key in ("tp", "fp", "fn"))
    result = {"selected_windows": len(frame), "scored_windows": n, "event_windows": e,
              "cow_groups": scored.cow_group.nunique(), "tp": tp, "fp": fp, "fn": fn,
              "event_f1": 2 * tp / (2 * tp + fp + fn) if e and 2 * tp + fp + fn else np.nan}
    if not n:
        return dict(result, rr_mae_bpm=np.nan, rr_rmse_bpm=np.nan, rr_r2=np.nan,
                    mean_count_accuracy=np.nan, exact_count=np.nan, exact_count_rate=np.nan, exact_display=None)
    y, p = scored.truth_rr_bpm.to_numpy(float), scored.predicted_rr_bpm.to_numpy(float)
    error = p-y
    sst = float(np.sum((y-y.mean()) ** 2))
    accuracy = [max(0., 1-abs(row.predicted_count-row.truth_count)/row.truth_count)
                if row.truth_count > 0 else float(row.predicted_count == 0) for row in scored.itertuples()]
    exact = int(scored.predicted_count.eq(scored.truth_count).sum())
    return dict(result, rr_mae_bpm=float(abs(error).mean()), rr_rmse_bpm=float(np.sqrt(np.mean(error**2))),
                rr_r2=float(1-np.sum(error**2)/sst) if sst else np.nan,
                mean_count_accuracy=float(np.mean(accuracy)), exact_count=exact,
                exact_count_rate=exact/n, exact_display=f"{exact}/{n}")


def paired_ci(candidate, control):
    left = candidate[candidate.eligible_rr].set_index("segment_id")
    right = control[control.eligible_rr].set_index("segment_id")
    assert left.index.is_unique and right.index.is_unique and set(left.index) == set(right.index)
    if not len(left):
        return None
    right = right.reindex(left.index)
    assert left.cow_group.eq(right.cow_group).all()
    delta = left.absolute_rr_error_bpm-right.absolute_rr_error_bpm
    groups = pd.DataFrame({"cow": left.cow_group, "delta": delta}).groupby("cow").delta.agg(["sum", "size"])
    draws = np.random.default_rng(20260929).integers(0, len(groups), size=(10000, len(groups)))
    boot = groups["sum"].to_numpy()[draws].sum(axis=1)/groups["size"].to_numpy()[draws].sum(axis=1)
    low, high = np.quantile(boot, [.025, .975])
    return {"windows": len(left), "cow_groups": len(groups), "delta_mae_bpm": float(delta.mean()),
            "ci95_low": float(low), "ci95_high": float(high), "draws": 10000, "seed": 20260929}


def validate_replay(path, verify_curves=False):
    prediction = pd.read_csv(path / "predictions.csv")
    registry = pd.read_csv(path / "reference_registry.csv")
    completed = json.loads((path / "COMPLETED.json").read_text(encoding="utf-8"))
    assert len(prediction) == completed["predictions"] == 796
    assert len(registry) == completed["segments"] == 199
    assert not prediction.duplicated(["segment_id", "arm"]).any() and registry.segment_id.is_unique
    assert prediction.groupby("segment_id").arm.agg(set).map(lambda value: value == set(ARMS)).all()
    assert set(prediction.segment_id) == set(registry.segment_id)
    historical = "grouped_split" if "grouped_split" in prediction else "split"
    prediction["historical_group"] = prediction[historical]
    assert np.allclose(prediction.predicted_rr_bpm, 60*prediction.predicted_count/prediction.duration_seconds)
    scored = prediction[prediction.eligible_rr]
    assert scored.segment_id.nunique() == 155
    assert np.allclose(scored.truth_rr_bpm, 60*scored.truth_count/scored.duration_seconds)
    assert np.allclose(scored.absolute_rr_error_bpm, abs(scored.predicted_rr_bpm-scored.truth_rr_bpm))
    accuracy = [max(0., 1-abs(row.predicted_count-row.truth_count)/row.truth_count)
                if row.truth_count > 0 else float(row.predicted_count == 0) for row in scored.itertuples()]
    assert np.allclose(accuracy, scored.mean_count_accuracy)
    assert np.array_equal(scored.exact_count, scored.predicted_count.eq(scored.truth_count).astype(int))
    assert prediction.loc[~prediction.eligible_rr, ["truth_rr_bpm", "absolute_rr_error_bpm", "mean_count_accuracy"]].isna().all().all()
    matches = pd.read_csv(path / "event_matches.csv")
    hits = matches[matches.type.eq("TP")]
    assert hits.error_seconds.abs().le(.30+1e-10).all()
    for key in ("predicted_index", "reference_index"):
        assert not hits.duplicated(["segment_id", "arm", key]).any()
    frequencies = matches.groupby(["segment_id", "arm", "type"]).size()
    for row in prediction[prediction.eligible_events].itertuples():
        for key in ("tp", "fp", "fn"):
            assert getattr(row, key) == frequencies.get((row.segment_id, row.arm, key.upper()), 0)
    metrics = pd.read_csv(path / "metrics.csv")
    for row in metrics.itertuples():
        pool = prediction[prediction.cohort.eq(row.cohort) & prediction.arm.eq(row.arm)]
        if row.split != "all":
            pool = pool[pool.split.eq(row.split)]
        rebuilt = measure(pool)
        for key in ("selected_windows", "scored_windows", "rr_mae_bpm", "rr_rmse_bpm", "rr_r2", "exact_count", "mean_count_accuracy", "event_f1", "tp", "fp", "fn"):
            same(rebuilt[key], getattr(row, key), f"{path.name}/{row.cohort}/{row.arm}/{key}")
    intervals = pd.read_csv(path / "paired_intervals.csv")
    for row in intervals.itertuples():
        pool = prediction[prediction.cohort.eq(row.cohort)]
        if row.split != "all":
            pool = pool[pool.split.eq(row.split)]
        rebuilt = paired_ci(pool[pool.arm.eq(row.candidate)], pool[pool.arm.eq(row.control)])
        for key in ("windows", "cow_groups", "delta_mae_bpm", "ci95_low", "ci95_high"):
            same(rebuilt[key], getattr(row, key), f"stored interval {key}")
    curve_count = 0
    if verify_curves:
        events = pd.read_csv(path / "predicted_events.csv")
        grouped = events.groupby(["segment_id", "arm"])
        for row in prediction.itertuples():
            curve = pd.read_csv(path / "curves" / row.arm / f"{row.segment_id}.csv")
            saved = grouped.get_group((row.segment_id, row.arm)) if (row.segment_id, row.arm) in grouped.groups else events.iloc[:0]
            peaks = curve[curve.is_peak]
            assert len(peaks) == len(saved) == row.predicted_count
            assert np.array_equal(peaks.frame_index.to_numpy(int), saved.frame_index.to_numpy(int))
            assert np.allclose(peaks.time_seconds, saved.time_seconds)
            curve_count += 1
        assert len(list((path / "temperatures").glob("*.csv"))) == 199
        assert len(list((path / "plots").glob("*.png"))) == 199
    return prediction, registry, {"metric_rows_recomputed": len(metrics), "stored_intervals_recomputed": len(intervals),
        "event_match_rows_checked": len(matches), "saved_peak_curves_checked": curve_count}


def best_metrics(training):
    import torch
    checkpoint = torch.load(training / "weights/best.pt", map_location="cpu", weights_only=False)
    epochs = pd.read_csv(training / "results.csv")
    assert epochs.epoch.tolist() == list(range(1, int(epochs.epoch.max())+1))
    mask = np.ones(len(epochs), dtype=bool)
    for key, value in checkpoint["train_metrics"].items():
        if key in epochs:
            mask &= np.isclose(epochs[key], value, rtol=0, atol=1e-10)
    selected = epochs[mask]
    assert len(selected) == 1, "Best checkpoint metric/epoch match must be unique"
    return selected.iloc[0], epochs


def wide_table(aggregated, historical="all"):
    rows = []
    for metric, label, unit in METRICS:
        for scope in SCOPES:
            subset = aggregated[aggregated.scope.eq(scope) & aggregated.historical_group.eq(historical)]
            sample = subset.iloc[0]
            scope_label = SCOPE_LABELS[scope]
            if historical != "all" and scope in ("all82", "all155"):
                prefix = "仅约30秒合计" if scope == "all82" else "全部合计"
                scope_label = f"{prefix}{int(sample.scored_windows)}（历史{'训练' if historical == 'train' else '验证'}组）"
            row = {"指标": label, "范围": scope_label, "RR评分窗": int(sample.scored_windows),
                   "事件评分窗": int(sample.event_windows), "纳入窗": int(sample.selected_windows), "单位或口径": unit}
            for model, prefix in (("grouped", "旧牛号"), ("random", "新随机")):
                for arm in ARMS:
                    selected = subset[subset.model.eq(model) & subset.arm.eq(arm)].iloc[0]
                    row[f"{prefix}_{arm}"] = selected[metric]
            rows.append(row)
    return pd.DataFrame(rows)


def main():
    if OUT.exists():
        raise FileExistsError(f"Preserve earlier exports: {OUT}")
    spec = importlib.util.spec_from_file_location("frame_random_runner_check", HERE / "run.py")
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    runner.verify_sources()
    sources = [HERE / "COMPLETED.json", HERE / "freeze.json", HERE / "split_audit.json", HERE / "comparison_metrics.csv", ROOT / "models/best.pt", Path(__file__)]
    for replay, training in ((NEW_REPLAY, NEW_TRAINING), (OLD_REPLAY, OLD_TRAINING)):
        sources += [replay / name for name in ("COMPLETED.json", "predictions.csv", "predicted_events.csv", "event_matches.csv",
            "metrics.csv", "paired_intervals.csv", "reference_registry.csv", "frame_times.csv", "freeze.json")]
        sources += [training / "results.csv", training / "weights/best.pt"]
    fingerprints = [{"path": str(path), "sha256": sha256(path)} for path in sources]
    current, registry, current_checks = validate_replay(NEW_REPLAY, True)
    previous, _, previous_checks = validate_replay(OLD_REPLAY)
    keys = ["segment_id", "arm"]
    paired = current.merge(previous, on=keys, suffixes=("_new", "_old"), validate="one_to_one")
    assert len(paired) == 796
    for key in ("truth_count", "duration_seconds", "eligible_rr", "eligible_events", "cow_group", "cohort", "historical_group"):
        a, b = paired[f"{key}_new"], paired[f"{key}_old"]
        assert (a.eq(b) | (a.isna() & b.isna())).all(), f"Reference/member changed: {key}"
    completion = json.loads((HERE / "COMPLETED.json").read_text(encoding="utf-8"))
    assert completion["weight_sha256"] == sha256(NEW_TRAINING / "weights/best.pt")
    for path, training in ((NEW_REPLAY, NEW_TRAINING), (OLD_REPLAY, OLD_TRAINING)):
        assert json.loads((path / "COMPLETED.json").read_text())["weights_sha256"] == sha256(training / "weights/best.pt")
    new_best, new_epochs = best_metrics(NEW_TRAINING)
    old_best, old_epochs = best_metrics(OLD_TRAINING)
    assert int(new_epochs.epoch.max()) == 200
    OUT.mkdir()
    aggregated = []
    for model, prediction in (("grouped", previous), ("random", current)):
        for historical in ("all", "train", "val"):
            for scope in SCOPES:
                for arm in ARMS:
                    pool = select(prediction[prediction.arm.eq(arm)], scope, historical)
                    aggregated.append({"model": model, "historical_group": historical, "scope": scope, "arm": arm, **measure(pool)})
    aggregate = pd.DataFrame(aggregated)
    save("rr_aggregates_long.csv", aggregate)
    stored_comparison = pd.read_csv(HERE / "comparison_metrics.csv")
    scope_map = {"primary47": "anchored49", "all82_approx30": "all82", "old73": "short73", "all155_mixed_duration": "all155", **{name: name for name in ("P1", "P2", "P3")}}
    for row in stored_comparison.itertuples():
        pool = aggregate[aggregate.scope.eq(scope_map[row.scope]) & aggregate.historical_group.eq("all") & aggregate.arm.eq(row.arm)].set_index("model")
        for source, model in (("old", "grouped"), ("new", "random")):
            for a, b in (("mae_bpm", "rr_mae_bpm"), ("rmse_bpm", "rr_rmse_bpm"), ("r2", "rr_r2"), ("mean_count_accuracy", "mean_count_accuracy"), ("exact_count", "exact_count"), ("event_f1", "event_f1")):
                same(pool.loc[model, b], getattr(row, f"{source}_{a}"), f"original comparison {row.scope}/{row.arm}/{a}")
    yolo = pd.DataFrame([{"指标": label, "旧牛号_鼻部框": old_best[f"metrics/{key}(B)"], "新随机_鼻部框": new_best[f"metrics/{key}(B)"],
        "旧牛号_鼻孔关键点": old_best[f"metrics/{key}(P)"], "新随机_鼻孔关键点": new_best[f"metrics/{key}(P)"],
        "说明": "旧val8316帧、新val8109帧，成员不同；均为best轮指标，非公平同测试集模型增益"}
        for key, label in (("precision", "Precision"), ("recall", "Recall"), ("mAP50", "mAP50"), ("mAP50-95", "mAP50-95"))])
    save("yolo_validation_best.csv", yolo)
    ci = []
    for scope in SCOPES:
        new_pool, old_pool = select(current, scope), select(previous, scope)
        for arm in ARMS:
            comparison = paired_ci(new_pool[new_pool.arm.eq(arm)], old_pool[old_pool.arm.eq(arm)])
            ci.append({"范围": SCOPE_LABELS[scope], "scope": scope, "比较": "新随机减旧牛号，同RR规则", "arm": arm, **comparison})
        for arm in ARMS[1:]:
            comparison = paired_ci(new_pool[new_pool.arm.eq(arm)], new_pool[new_pool.arm.eq("C0P0")])
            ci.append({"范围": SCOPE_LABELS[scope], "scope": scope, "比较": "新随机模型内候选减C0P0，历史补齐差异保留", "arm": arm, **comparison})
    intervals = pd.DataFrame(ci)
    save("paired_intervals_exploratory.csv", intervals)
    notes = [
        ["完成状态", "2026-10-21 05:05全部流水线完成；200轮上限跑满，非早停；199片段/796条四组预测/199张曲线图。"],
        ["训练及最佳轮", f"旧牛号训练最后{int(old_epochs.epoch.max())}轮、best第{int(old_best.epoch)}轮；新随机最后{int(new_epochs.epoch.max())}轮、best第{int(new_best.epoch)}轮。"],
        ["四组定义", "; ".join(f"{key}={value}" for key, value in LABELS.items())],
        ["表格格式", "每个方法独立一列；全部页每项按旧73、完整47、P1、P2、P3、仅约30秒82、全部155排序，155位于每项最后一行。历史训练/验证组页对应末行127/28、约30秒69/13，均非新随机帧验证集。82包含47/P1/P2/P3，不能再累加。"],
        ["快照类型", "按之前格式导入已冻结CSV计算结果的只读Excel快照，不含公式；重新计算由导出脚本完成，不是动态Excel模型。"],
        ["窗口与可评分", "199纳入，155有完整次数参考：旧73、47/49、5/8、6/16、24/53。排除44不完整参考及两段事后36次，参考、时长与成员未改。"],
        ["时长", "155是混合时长；仅82是约30秒完整参考。旧73沿冻结原时长，不把旧次数按30秒重算。RR=60N/T，预测与参考共用冻结T。"],
        ["事件参考", "仅82窗有完整事件参考；155行事件F1也只来自其中82窗。旧73F1/TP/FP/FN为空，不填0。"],
        ["平均相对计数准确度", "逐窗max(0,1-|预测次数-参考次数|/参考次数)平均，参考0时比较预测也为0；不是逐事件准确率或完全计数率。"],
        ["R²与F1", "R²=1-SSE/SST从合计逐窗重算，不平均各队列R²；F1由合计TP/FP/FN计算，不平均各窗F1。"],
        ["随机帧验证", "train32433/val8109，所有199片段及145牛号组跨train/val，31组相同编码图像哈希跨集合。完整RR窗含训练帧，无独立测试。"],
        ["历史分组诊断", "历史train/val页沿07牛号组仅供同窗追溯；新随机模型在这些窗也训练过。13窗不是本次随机帧未见RR验证，不能把8109离散验证帧当完整30秒测试窗。"],
        ["训练策略比较", "旧/新YOLO验证成员不同，不能用验证mAP差当公平模型增益。RR模型策略对照用相同参考/时长/窗口，但仍属暴露数据的开发比较。"],
        ["质量选择归因", "沿用原四组C1/C0时补齐不同，不能把包差归因于纯质量选择，不推翻08同补齐归因的不确定结论。"],
        ["配对区间", "全部探索性，牛号组10000次bootstrap、seed20260929、未多重校正；差值<0有利，不作为新的确认试验。"],
        ["采用决定", "此次只报告结果，不替换默认权重或RR规则；新旧各有改善和退化，不称所有指标或所有队列均提升。"],
        ["最佳新权重", str(NEW_TRAINING / "weights/best.pt")], ["最佳新权重SHA256", completion["weight_sha256"]],
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    tables = {}

    def sheet(name, frame, source):
        tables[name] = frame
        ws = add_sheet(workbook, name, frame, source, {"说明": 115, "范围": 31, "单位或口径": 30, "指标": 26, "项目": 26})
        if "指标" in frame and name.startswith("RR"):
            numeric_cols = [index+1 for index, column in enumerate(frame.columns) if column.startswith(("旧牛号_", "新随机_"))]
            for number, record in enumerate(frame.to_dict("records"), 2):
                rate = record["指标"] in ("平均相对计数准确度", "完全计数率")
                count = record["指标"] in ("事件TP", "事件FP", "事件FN")
                for index in numeric_cols:
                    ws.cell(number, index).number_format = "0.00%" if rate else "0" if count else "0.000000"
                if record.get("范围", "").startswith("全部合计"):
                    for cell in ws[number]:
                        cell.fill = PatternFill("solid", fgColor="DCEBE5")
                        cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
        return ws

    sheet("说明", pd.DataFrame(notes, columns=["项目", "说明"]), HERE / "PROTOCOL.md")
    ws = sheet("YOLO最佳轮验证", yolo, NEW_TRAINING / "results.csv")
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = "0.00%"
    for historical, name in (("all", "RR新旧8列"), ("train", "RR历史训练组_非新训练窗"), ("val", "RR历史验证组_非新验证窗")):
        wide = wide_table(aggregate, historical)
        save(f"rr_metrics_{historical}_wide.csv", wide)
        sheet(name, wide, NEW_REPLAY / "predictions.csv")
        if historical == "all":
            for prefix, title in (("新随机_", "RR新随机4列"), ("旧牛号_", "RR旧牛号4列")):
                columns = list(wide.columns[:6])+[column for column in wide.columns if column.startswith(prefix)]
                sheet(title, wide[columns], NEW_REPLAY / "predictions.csv" if prefix.startswith("新") else OLD_REPLAY / "predictions.csv")
    sheet("探索性配对区间", intervals, "paired absolute RR error; cow-group bootstrap")
    sheet("重算汇总长表", aggregate, NEW_REPLAY / "predictions.csv")
    sheet("流水线已有新旧比较", stored_comparison, HERE / "comparison_metrics.csv")
    sheet("新逐窗四组记录", current, NEW_REPLAY / "predictions.csv")
    sheet("旧逐窗四组记录", previous, OLD_REPLAY / "predictions.csv")
    sheet("参考及暴露", registry, NEW_REPLAY / "reference_registry.csv")
    sheet("新逐轮训练", new_epochs, NEW_TRAINING / "results.csv")
    sheet("来源哈希", pd.DataFrame(fingerprints), "SHA256 imported source files")
    workbook.save(OUT / "metric_comparison.xlsx")
    report = ["# 全帧随机8:2训练完成后的指标对照", "", "2026-10-21 05:05流水线完整结束。新模型跑满200轮，best第200轮；不是早停。", "",
              "## 定位验证", markdown_table(yolo), "", "旧/新定位验证成员不同，不能把此表差异称同测试集模型增益。", "", "## 四组方法",
              *[f"- {key}：{value}" for key, value in LABELS.items()], "", "## 新模型RR：方法独立成列"]
    for scope in ("anchored49", "all82", "all155", "short73", "P1", "P2", "P3"):
        pool = aggregate[aggregate.scope.eq(scope) & aggregate.model.eq("random") & aggregate.historical_group.eq("all")].set_index("arm")
        records = []
        for key, label, _ in METRICS:
            record = {"指标": label}
            record.update({arm: f"{pool.loc[arm,key]:.2%}" if key in ("mean_count_accuracy", "exact_count_rate") else pool.loc[arm,key] for arm in ARMS})
            records.append(record)
        report += ["", f"### {SCOPE_LABELS[scope]}（RR {int(pool.iloc[0].scored_windows)}窗，事件{int(pool.iloc[0].event_windows)}窗）", markdown_table(pd.DataFrame(records))]
    report += ["", "## 同规则模型策略对照", markdown_table(stored_comparison), "", "## 探索性配对区间", markdown_table(intervals),
               "", "## 解释边界", *[f"- {key}：{value}" for key, value in notes[3:16]], "", "## 核验范围",
               "由逐窗记录重算全部聚合；重建原汇总/区间，核对796份曲线峰标记与事件，回读全部表格数值和每项最后合计行。源文件前后SHA不变。",
               "工作簿为只读数值快照，无公式/错误单元格；未做Excel界面渲染。没有重训、改参考/窗口/方法，未将该检查称为完整独立实验审计。"]
    (OUT / "指标对照说明.md").write_text("\n".join(report)+"\n", encoding="utf-8")
    reopened = load_workbook(OUT / "metric_comparison.xlsx", data_only=False)
    checked_cells = 0
    for name, frame in tables.items():
        ws = reopened[name]
        assert ws.max_row == len(frame)+1 and ws.max_column == len(frame.columns)
        for r, row in enumerate(frame.itertuples(index=False, name=None), 2):
            for c, expected in enumerate(row, 1):
                actual = ws.cell(r, c).value
                expected = excel_value(expected)
                if isinstance(expected, (int, float)) and not isinstance(expected, bool):
                    same(actual, expected, f"Workbook {name}/{r}/{c}")
                else:
                    assert actual == expected, f"Workbook {name}/{r}/{c}: {actual!r}!={expected!r}"
                assert ws.cell(r,c).data_type not in ("f", "e")
                checked_cells += 1
        if name.startswith("RR"):
            labels = frame["指标"].drop_duplicates()
            for label in labels:
                final = frame[frame["指标"].eq(label)].iloc[-1]
                expected_n = 127 if name == "RR历史训练组_非新训练窗" else 28 if name == "RR历史验证组_非新验证窗" else 155
                assert final["范围"].startswith(f"全部合计{expected_n}") and final["RR评分窗"] == expected_n
    for item in fingerprints:
        assert sha256(Path(item["path"])) == item["sha256"], f"Imported source changed: {item['path']}"
    runner.verify_sources()
    checks = {"status": "passed", "date": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(), "new_replay": current_checks,
        "old_replay": previous_checks, "aggregate_rows": len(aggregate), "existing_comparison_rows_checked": len(stored_comparison),
        "exploratory_paired_intervals": len(intervals), "new_best_epoch": int(new_best.epoch), "new_last_epoch": int(new_epochs.epoch.max()),
        "old_best_epoch": int(old_best.epoch), "old_last_epoch": int(old_epochs.epoch.max()), "workbook_sheets": reopened.sheetnames,
        "workbook_cells_roundtripped": checked_cells, "last_row_all155_historical_train127_val28_checked": True, "formulas": 0, "error_cells": 0,
        "missing_metrics_preserved": True, "source_files_unchanged": True, "excel_ui_render": "not run", "default_changed": False}
    json_file("verification.json", checks)
    json_file("sources.json", fingerprints)
    json_file("COMPLETED.json", {"status": "complete", "scope": "completed-run metric export, not independent confirmation", "verification": "verification.json"})
    print(json.dumps(checks, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
