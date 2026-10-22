"""Exploratory pooled summaries; keep original cohorts and duration provenance."""

import json
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill


ARMS = ("C0P0", "C1P0", "C0P1", "C1P1")
COHORTS = ("short73", "anchored49", "P1", "P2", "P3")
CONTRASTS = (("C1P0", "C0P0"), ("C1P1", "C0P1"), ("C0P1", "C0P0"),
             ("C1P1", "C1P0"), ("C1P1", "C0P0"))
DISPLAY = {"short73": "本次实际旧73输入（非新30秒版）", "anchored49": "49窗完整参考", "P1": "P1", "P2": "P2", "P3": "P3"}
METRICS = (("rr_mae_bpm", "RR MAE", "0.000000"), ("rr_rmse_bpm", "RR RMSE", "0.000000"),
           ("rr_r2", "RR R²", "0.000000"), ("mean_count_accuracy", "RR平均相对准确度", "0.00%"),
           ("exact_count_rate", "完全计数率", "0.00%"), ("exact_count_display", "完全计数/可评分", "@"),
           ("event_f1", "事件F1", "0.000000"), ("tp", "事件TP", "0"),
           ("fp", "事件FP", "0"), ("fn", "事件FN", "0"))


def aggregate(predictions, scope="pooled_mixed_duration"):
    rows = []
    for split in ("all", "train", "val"):
        pool = predictions if split == "all" else predictions[predictions.split.eq(split)]
        for arm in ARMS:
            selected = pool[pool.arm.eq(arm)]
            group = selected[selected.eligible_rr]
            if group.empty:
                continue
            truth = group.truth_rr_bpm.to_numpy(float)
            error = group.predicted_rr_bpm.to_numpy(float) - truth
            sst = np.sum((truth - truth.mean()) ** 2)
            events = group[group.eligible_events]
            tp, fp, fn = (int(events[key].sum()) for key in ("tp", "fp", "fn"))
            rows.append({"scope": scope, "cohort": scope, "split": split, "arm": arm,
                         "scored_windows": len(group), "selected_windows": len(selected),
                         "cow_groups": group.cow_group.nunique(), "event_scored_windows": len(events),
                         "rr_mae_bpm": float(np.mean(np.abs(error))),
                         "rr_rmse_bpm": float(np.sqrt(np.mean(error ** 2))),
                         "rr_r2": float(1 - np.sum(error ** 2) / sst) if sst > 0 else np.nan,
                         "exact_count": int(group.exact_count.sum()),
                         "exact_count_rate": float(group.exact_count.mean()),
                         "mean_count_accuracy": float(group.mean_count_accuracy.mean()),
                         "tp": tp if len(events) else np.nan, "fp": fp if len(events) else np.nan,
                         "fn": fn if len(events) else np.nan,
                         "event_f1": 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else np.nan})
    return pd.DataFrame(rows)


def paired(predictions, scope="pooled_mixed_duration", draws=10000, seed=20260929):
    rows = []
    eligible = predictions[predictions.eligible_rr]
    for split in ("all", "train", "val"):
        pool = eligible if split == "all" else eligible[eligible.split.eq(split)]
        for candidate, control in CONTRASTS:
            base = pool[pool.arm.eq(control)].set_index("segment_id")
            other = pool[pool.arm.eq(candidate)].set_index("segment_id").reindex(base.index)
            if base.empty:
                continue
            assert other.absolute_rr_error_bpm.notna().all()
            delta = other.absolute_rr_error_bpm - base.absolute_rr_error_bpm
            clusters = pd.DataFrame({"cow_group": base.cow_group, "delta": delta}).groupby("cow_group").delta.agg(["sum", "size"])
            rng = np.random.default_rng(seed)
            sample = rng.integers(0, len(clusters), size=(draws, len(clusters)))
            boot = clusters["sum"].to_numpy()[sample].sum(axis=1) / clusters["size"].to_numpy()[sample].sum(axis=1)
            low, high = np.quantile(boot, [.025, .975])
            rows.append({"scope": scope, "split": split, "candidate": candidate, "control": control,
                         "windows": len(base), "cow_groups": len(clusters), "delta_mae_bpm": float(delta.mean()),
                         "ci95_low": float(low), "ci95_high": float(high),
                         "relative_mae_reduction": float(-delta.mean() / base.absolute_rr_error_bpm.mean())
                         if base.absolute_rr_error_bpm.mean() > 0 else np.nan,
                         "draws": draws, "seed": seed,
                         "interpretation": "Post-hoc development aggregation; unadjusted exploratory CI, not independent confirmation"})
    return pd.DataFrame(rows)


def audit_73(root, replay, registry, sha256):
    import cv2

    segments = pd.read_csv(root / "corpus_v1" / "segments.csv")
    frames = pd.read_csv(replay / "frames.csv")
    times = pd.read_csv(replay / "frame_times.csv")
    fixed_root = root.parents[1] / "Dataset_new" / "72video" / "lindian_fixed30_frames"
    rows = []
    reference = registry.set_index("segment_id")
    for segment in segments[segments.cohort.eq("short73")].itertuples(index=False):
        source = Path(segment.source_path)
        cap = cv2.VideoCapture(str(source))
        opened = cap.isOpened()
        count = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = cap.get(cv2.CAP_PROP_FPS)
        cap.release()
        used = frames[frames.segment_id.eq(segment.segment_id)]
        binding = times[times.segment_id.eq(segment.segment_id)]
        video_id = segment.segment_id.removeprefix("S73_")
        candidate = fixed_root / video_id
        candidate_frames = list(candidate.glob("frame_*.jpg")) if candidate.is_dir() else []
        rows.append({"segment_id": segment.segment_id, "source_path": str(source), "source_opened": opened,
                     "source_sha256_matches_frozen": sha256(source) == segment.source_sha256,
                     "source_frame_count": count, "source_fps": fps,
                     "source_container_duration_seconds": count / fps if fps > 0 else np.nan,
                     "replayed_frame_count": len(used),
                     "replayed_duration_seconds": reference.loc[segment.segment_id, "duration_seconds"],
                     "first_replayed_image": used.iloc[0].image_path,
                     "last_bound_time_seconds": binding.time_seconds.max(),
                     "other_30s_candidate_dir": str(candidate) if candidate.is_dir() else "",
                     "other_30s_candidate_frames": len(candidate_frames),
                     "candidate_used_by_this_replay": False,
                     "note": "User reports new 30s replacements; actual replay still uses legacy frames/reference. No transfer of counts or labels."})
    return pd.DataFrame(rows)


def reconcile_reviewed_inputs(root, replay, sha256):
    reviews = pd.read_csv(root / "corpus_v1" / "manual_review_status.csv")
    frames = pd.read_csv(replay / "frames.csv")
    active = reviews[reviews.cohort.ne("R3_jiufu")]
    new = frames[frames.cohort.ne("short73")]
    keys = new.segment_id + ":" + new.source_frame.map(lambda value: f"{int(value):06d}")
    equal = set(active.frame_key) == set(keys)
    assert equal and active.review_status.eq("approved").all()
    anchored = root.parents[1] / "Dataset_new" / "72video" / "lindian_anchored30_frames"
    used = set(new.loc[new.cohort.eq("anchored49"), "segment_id"].str.removeprefix("A49_"))
    directories = {path.name for path in anchored.iterdir() if path.is_dir()}
    example = "A49_16170075"
    chosen = new[new.segment_id.eq(example)].iloc[0]
    materialized = root / "yolo_dataset_lindian_20261016_v1" / "images" / chosen.split / "A49_16170075__000000.jpg"
    identical_image = sha256(materialized) == sha256(Path(chosen.image_path))
    materialized_label = root / "yolo_dataset_lindian_20261016_v1" / "labels" / chosen.split / "A49_16170075__000000.txt"
    manual_label = Path(active[active.frame_key.eq("A49_16170075:000000")].iloc[0].manual_label_path)
    identical_label = sha256(materialized_label) == sha256(manual_label)
    assert identical_image and identical_label
    ppt = Path("E:/real/汇报+论文/组会/第三次工作汇报材料/cow_rr_group_meeting_20260908_v4.pptx")
    ppt_evidence = {}
    if ppt.is_file():
        with ZipFile(ppt) as archive:
            for number in (4, 12, 16):
                xml = ElementTree.fromstring(archive.read(f"ppt/slides/slide{number}.xml"))
                ppt_evidence[str(number)] = " | ".join(node.text or "" for node in xml.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"))
    return {"user_supplied_review_table": str(root / "corpus_v1" / "manual_review_status.csv"),
            "reviewed_new_frames": len(active), "reviewed_new_segment_count": active.frame_key.str.split(":").str[0].nunique(),
            "approved_new_frames_equal_replayed_new_frame_keys": equal,
            "review_frame_counts_by_cohort": {key: int(value) for key, value in active.groupby("cohort").size().items()},
            "actual_replayed_frame_counts_by_cohort": {key: int(value) for key, value in frames.groupby("cohort").size().items()},
            "anchored_frame_directories": len(directories), "anchored_replayed_segments": len(used),
            "anchored_directories_not_in_replay": sorted(directories - used),
            "example": {"old_id": "S73_16170075", "old_frames": int(frames.segment_id.eq("S73_16170075").sum()),
                        "new_30s_id": example, "new_30s_frames": int(new.segment_id.eq(example).sum()),
                        "review_json_count": len(list((root / "biaozhu" / "16170075").glob("*.json"))),
                        "materialized_training_image_matches_source": identical_image,
                        "materialized_training_label_matches_reviewed_label": identical_label},
            "ppt_path": str(ppt), "ppt_sha256": sha256(ppt) if ppt.is_file() else None,
            "ppt_slide_text_evidence": ppt_evidence,
            "interpretation": "Reviewed 30s frames and labels were used; pipeline also retains old73 as a separate cohort. Do not infer source from R2."}


def augment(workbook, report, predictions, metrics, registry, replay, output, root,
            add_sheet, markdown_table, sha256):
    pooled = aggregate(predictions)
    approx30 = aggregate(predictions[predictions.cohort.ne("short73")], "approx30_only")
    intervals = pd.concat([paired(predictions), paired(predictions[predictions.cohort.ne("short73")], "approx30_only")], ignore_index=True)
    pooled.to_csv(output / "pooled_metrics.csv", index=False, encoding="utf-8-sig")
    approx30.to_csv(output / "approx30_only_metrics.csv", index=False, encoding="utf-8-sig")
    intervals.to_csv(output / "pooled_paired_intervals.csv", index=False, encoding="utf-8-sig")
    audit = audit_73(root, replay, registry, sha256)
    audit.to_csv(output / "73_input_duration_audit.csv", index=False, encoding="utf-8-sig")
    reconciliation = reconcile_reviewed_inputs(root, replay, sha256)
    (output / "input_membership_reconciliation.json").write_text(json.dumps(reconciliation, ensure_ascii=False, indent=2), encoding="utf-8")
    warning = ("人工复核30秒数据已用于训练和复放：anchored49共12741帧，另含P1/P2/P3；"
               "流程同时保留旧73共7488帧及6–20秒参考。因此155条是混合时长探索性汇总，不是155个独立新30秒窗；"
               "只计本次约30秒完整参考时为82窗，验证组为13窗。")
    workbook["说明"].append(["本轮输入更正警示", warning])
    workbook["说明"].append(["合计方式", "按逐窗记录重算；MAE按窗数加权，RMSE按平方误差重算，R²由全部真值/预测重算，F1只汇总有事件参考的82窗。"])
    workbook["说明"].append(["3.790651的含义", "是五个队列MAE的等权宏平均；不是155条记录的MAE（对照为3.916826）。两个量不能混名。"])
    leading = ["# 本次实际输入与汇总补充\n", warning,
               "\n## 人工复核是否实际使用\n",
               "已逐键核对：复核表中林甸33054个approved新帧，与本次复放的全部非旧73帧完全对应。"
               "16170075的30秒帧260张、JSON260份，训练副本图像和标签与相应来源哈希一致；"
               "另有旧版124帧，保存在S73而非A49组。不是未使用你的标注。\n",
               "给定anchored目录实际有50个子目录，本次49个纳入，额外bs1479未纳入；"
               "不能把它直接说成73个均已进入复放的新30秒窗口。"
               "已读取指定PPT：第4页明确列73为6–20秒、49为约30秒；第12页历史RR R²为0.946595，"
               "不与本次新YOLO、固定四组对照混作一项。低R²本身不能反推出使用了哪套数据。\n",
               "\n**注意：下面155/127/28为当前结果记录数，不代表不同牛只数。RR可评分155窗来自105牛号组；验证28窗来自17牛号组。**\n",
               "准确度同时列出平均相对准确度与完全计数率：前者不是分类正确率，也不是事件F1。\n",
               "## 合计的正确计算\n",
               "MAE使用全部逐窗绝对误差平均，不能将5个队列MAE简单平均后标为155窗；R²不能平均。"
               "3.790651是五队列等权MAE，而当前155窗对照MAE是3.916826。事件指标合计实际仅有82窗参考。\n"]
    for split in ("all", "train", "val"):
        total = pooled[pooled.split.eq(split)].set_index("arm")
        rows = []
        for key, label, _ in METRICS:
            for cohort in COHORTS + ("pooled_mixed_duration",):
                block = total if cohort == "pooled_mixed_duration" else metrics[metrics.cohort.eq(cohort) & metrics.split.eq(split)].set_index("arm")
                rr_count = int(block.iloc[0].scored_windows) if not block.empty else 0
                pool = registry[registry.eligible_rr]
                if split != "all":
                    pool = pool[pool.split.eq(split)]
                if cohort != "pooled_mixed_duration":
                    pool = pool[pool.cohort.eq(cohort)]
                event_count = int(pool.eligible_events.sum())
                row = {"指标": label, "metric_key": key,
                       "队列": f"合计{rr_count}（混合时长，探索性）" if cohort == "pooled_mixed_duration" else DISPLAY[cohort],
                       "RR可评分窗": rr_count, "事件可评分窗": event_count}
                for arm in ARMS:
                    if block.empty:
                        value = None
                    elif key == "exact_count_display":
                        value = f"{int(block.loc[arm, 'exact_count'])}/{rr_count}"
                    elif key == "exact_count_rate":
                        value = float(block.loc[arm, "exact_count"] / rr_count)
                    else:
                        value = block.loc[arm, key]
                    row[arm] = value
                row["备注"] = warning if cohort == "pooled_mixed_duration" else "原参考和时长未改；空值不填0"
                rows.append(row)
        frame = pd.DataFrame(rows)
        frame.to_csv(output / f"by_metric_{split}_with_totals.csv", index=False, encoding="utf-8-sig")
        sheet = add_sheet(workbook, f"按指标_{split}", frame.drop(columns="metric_key"), output / "pooled_metrics.csv",
                          {"指标": 27, "队列": 43, "RR可评分窗": 15, "事件可评分窗": 15, "备注": 68})
        for number, row in enumerate(frame.to_dict("records"), 2):
            fmt = next(item[2] for item in METRICS if item[0] == row["metric_key"])
            for arm in ARMS:
                sheet.cell(number, list(frame.drop(columns="metric_key").columns).index(arm) + 1).number_format = fmt
            if row["队列"].startswith("合计"):
                for cell in sheet[number]:
                    cell.fill = PatternFill("solid", fgColor="DCEBE6")
                    cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="173B31")
        leading.append(f"\n## {split}：每项指标末行是合计\n")
        for key, label, fmt in METRICS:
            table = frame[frame.metric_key.eq(key)][["队列", "RR可评分窗", "事件可评分窗", *ARMS]].astype(object)
            if fmt == "0.00%":
                for arm in ARMS:
                    table[arm] = table[arm].map(lambda value: f"{value:.2%}" if pd.notna(value) else None)
            leading += [f"\n### {label}\n", markdown_table(table)]
    add_sheet(workbook, "合计_混合时长", pooled, replay / "predictions.csv", {"scope": 30, "cohort": 30})
    add_sheet(workbook, "仅约30秒窗", approx30, replay / "predictions.csv", {"scope": 30, "cohort": 30})
    add_sheet(workbook, "合计配对区间", intervals, replay / "predictions.csv", {"scope": 30, "interpretation": 80})
    add_sheet(workbook, "73输入时长核查", audit, root / "corpus_v1" / "segments.csv", {"source_path": 96, "note": 98})
    for name in ("按指标_val", "按指标_train", "按指标_all"):
        workbook.move_sheet(workbook[name], offset=-workbook.index(workbook[name]))
    workbook.active = 0
    leading += ["\n## 探索性配对区间\n", markdown_table(intervals.drop(columns="interpretation")),
                "\n后续原始分队列表仍保留，不能把输入核查警示理解为已经重跑了新的73个30秒窗。\n"]
    report[:0] = leading
    summary = {"warning": warning, "pooled": pooled.to_dict("records"),
               "approx30_only": approx30.to_dict("records"), "paired_intervals": intervals.to_dict("records"),
               "reviewed_30s_frames_were_used": True,
               "audit_73": {"rows": len(audit), "replayed_frames": int(audit.replayed_frame_count.sum()),
                            "duration_min_seconds": float(audit.replayed_duration_seconds.min()),
                            "duration_max_seconds": float(audit.replayed_duration_seconds.max()),
                            "all_source_hashes_match": bool(audit.source_sha256_matches_frozen.all()),
                            "all_video_metadata_frame_counts_match": bool(audit.source_frame_count.eq(audit.replayed_frame_count).all())}}
    (output / "aggregation_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    return {"pooled_metric_rows": len(pooled), "approx30_metric_rows": len(approx30),
            "pooled_paired_interval_rows": len(intervals), "audited_73_sources": len(audit),
            "pooled_all_rr_windows": int(pooled[pooled.split.eq("all")].scored_windows.iloc[0]),
            "pooled_val_rr_windows": int(pooled[pooled.split.eq("val")].scored_windows.iloc[0]),
            "pooling_status": "exploratory_mixed_duration_not_155_new_30s", "audit_73": summary["audit_73"]}
