"""Export checked, read-only metric snapshots with one column per RR arm."""

import argparse
import hashlib
import json
import math
import os
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
ARMS = ("C0P0", "C1P0", "C0P1", "C1P1")
COHORTS = ("short73", "anchored49", "P1", "P2", "P3")
SPLITS = ("all", "train", "val")
ARM_LABELS = {
    "C0P0": "原文融合 + 原文峰规则",
    "C1P0": "质量选择 + 原文峰规则",
    "C0P1": "原文融合 + 约束峰规则",
    "C1P1": "质量选择 + 约束峰规则",
}
COHORT_LABELS = {"short73": "B73-I短片", "anchored49": "49窗中的完整参考", "P1": "P1", "P2": "P2", "P3": "P3"}
SPLIT_LABELS = {"all": "全部窗口", "train": "YOLO训练牛号组", "val": "YOLO验证牛号组"}
METRICS = (
    ("scored_windows", "RR可评分窗口数", "窗", "0"),
    ("selected_windows", "纳入窗口数", "窗", "0"),
    ("rr_mae_bpm", "RR MAE", "次/分，越低越好", "0.000000"),
    ("rr_rmse_bpm", "RR RMSE", "次/分，越低越好", "0.000000"),
    ("rr_r2", "RR R²", "越高越好", "0.000000"),
    ("exact_count", "完全计数相同窗口数", "窗", "0"),
    ("exact_count_display", "完全计数相同/可评分", "窗/窗，不是事件准确率", "@"),
    ("mean_count_accuracy", "平均计数准确度", "百分比，越高越好", "0.00%"),
    ("tp", "事件TP", "次", "0"),
    ("fp", "事件FP", "次", "0"),
    ("fn", "事件FN", "次", "0"),
    ("event_f1", "事件F1", "0至1，±0.30秒匹配，越高越好", "0.000000"),
)


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def same(actual, expected, context):
    if not np.isclose(actual, expected, rtol=1e-10, atol=1e-10, equal_nan=True):
        raise AssertionError(f"{context}: {actual!r} != {expected!r}")


def check_sources(replay, training, dataset):
    os.environ.setdefault("YOLO_CONFIG_DIR", str(ROOT / "ultralytics_config"))
    import torch

    predictions = pd.read_csv(replay / "predictions.csv")
    metrics = pd.read_csv(replay / "metrics.csv")
    intervals = pd.read_csv(replay / "paired_intervals.csv")
    registry = pd.read_csv(replay / "reference_registry.csv")
    events = pd.read_csv(replay / "event_matches.csv")
    epochs = pd.read_csv(training / "results.csv")
    completed = json.loads((replay / "COMPLETED.json").read_text(encoding="utf-8"))
    ready = json.loads((dataset / "READY.json").read_text(encoding="utf-8"))
    assert predictions.arm.isin(ARMS).all()
    assert predictions.groupby("segment_id").arm.agg(set).map(lambda value: value == set(ARMS)).all()
    assert not predictions.duplicated(["segment_id", "arm"]).any()
    assert not registry.segment_id.duplicated().any()
    assert set(predictions.segment_id) == set(registry.segment_id)
    assert len(predictions) == completed["predictions"] == 4 * len(registry)
    assert len(registry) == completed["segments"]
    assert len(epochs) > 0 and epochs.epoch.tolist() == list(range(1, int(epochs.epoch.max()) + 1))
    assert not metrics.duplicated(["cohort", "split", "arm"]).any()
    for name in registry.columns:
        if name == "segment_id":
            continue
        expected = registry.set_index("segment_id")[name]
        actual = predictions.drop_duplicates("segment_id").set_index("segment_id")[name]
        pd.testing.assert_series_equal(actual.sort_index(), expected.sort_index())
        assert predictions.groupby("segment_id")[name].nunique(dropna=False).eq(1).all()
    assert np.allclose(predictions.predicted_rr_bpm,
                       60 * predictions.predicted_count / predictions.duration_seconds)
    valid = predictions[predictions.eligible_rr]
    assert np.allclose(valid.truth_rr_bpm, 60 * valid.truth_count / valid.duration_seconds)
    assert np.allclose(valid.absolute_rr_error_bpm, abs(valid.predicted_rr_bpm - valid.truth_rr_bpm))
    expected_accuracy = np.array([
        max(0., 1 - abs(row.predicted_count - row.truth_count) / row.truth_count)
        if row.truth_count > 0 else float(row.predicted_count == 0)
        for row in valid.itertuples(index=False)
    ])
    assert np.allclose(valid.mean_count_accuracy, expected_accuracy)
    assert np.array_equal(valid.exact_count, valid.predicted_count.eq(valid.truth_count).astype(int))
    assert predictions.loc[~predictions.eligible_rr,
                           ["truth_rr_bpm", "absolute_rr_error_bpm", "exact_count", "mean_count_accuracy"]].isna().all().all()
    assert events.type.isin(["TP", "FP", "FN"]).all()
    hits = events[events.type.eq("TP")]
    assert hits.error_seconds.abs().le(.30 + 1e-10).all()
    assert not hits.duplicated(["segment_id", "arm", "predicted_index"]).any()
    assert not hits.duplicated(["segment_id", "arm", "reference_index"]).any()
    counts = events.groupby(["segment_id", "arm", "type"]).size()
    for row in predictions[predictions.eligible_events].itertuples(index=False):
        for kind in ("TP", "FP", "FN"):
            assert int(getattr(row, kind.lower())) == counts.get((row.segment_id, row.arm, kind), 0)
    for row in metrics.itertuples(index=False):
        pool = predictions[predictions.cohort.eq(row.cohort) & predictions.arm.eq(row.arm)]
        if row.split != "all":
            pool = pool[pool.split.eq(row.split)]
        scored = pool[pool.eligible_rr]
        assert row.selected_windows == len(pool) and row.scored_windows == len(scored)
        errors = scored.predicted_rr_bpm.to_numpy() - scored.truth_rr_bpm.to_numpy()
        total = np.sum((scored.truth_rr_bpm - scored.truth_rr_bpm.mean()) ** 2)
        same(np.abs(errors).mean(), row.rr_mae_bpm, "MAE")
        same(np.sqrt(np.mean(errors ** 2)), row.rr_rmse_bpm, "RMSE")
        same(1 - np.sum(errors ** 2) / total if total > 0 else np.nan, row.rr_r2, "R2")
        same(scored.exact_count.sum(), row.exact_count, "exact count")
        same(scored.mean_count_accuracy.mean(), row.mean_count_accuracy, "count accuracy")
        event_rows = scored[scored.eligible_events]
        for key in ("tp", "fp", "fn"):
            same(event_rows[key].sum() if len(event_rows) else np.nan, getattr(row, key), key)
        tp, fp, fn = (event_rows[key].sum() for key in ("tp", "fp", "fn"))
        same(2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else np.nan, row.event_f1, "event F1")
    for row in intervals.itertuples(index=False):
        pool = predictions[predictions.cohort.eq(row.cohort) & predictions.eligible_rr]
        if row.split != "all":
            pool = pool[pool.split.eq(row.split)]
        control = pool[pool.arm.eq(row.control)].set_index("segment_id")
        candidate = pool[pool.arm.eq(row.candidate)].set_index("segment_id").reindex(control.index)
        delta = candidate.absolute_rr_error_bpm - control.absolute_rr_error_bpm
        clusters = pd.DataFrame({"cow": control.cow_group, "delta": delta}).groupby("cow").delta.agg(["sum", "size"])
        rng = np.random.default_rng(int(row.seed))
        draws = rng.integers(0, len(clusters), size=(int(row.draws), len(clusters)))
        boot = clusters["sum"].to_numpy()[draws].sum(axis=1) / clusters["size"].to_numpy()[draws].sum(axis=1)
        low, high = np.quantile(boot, [.025, .975])
        assert row.windows == len(control) and row.cow_groups == len(clusters)
        same(delta.mean(), row.delta_mae_bpm, "paired delta")
        same(low, row.ci95_low, "CI low")
        same(high, row.ci95_high, "CI high")
    checkpoint = torch.load(training / "weights" / "best.pt", map_location="cpu", weights_only=False)
    stored = checkpoint["train_metrics"]
    mask = np.ones(len(epochs), dtype=bool)
    for key, value in stored.items():
        if key in epochs:
            mask &= np.isclose(epochs[key], value, rtol=0., atol=1e-10)
    matched = epochs.loc[mask]
    assert len(matched) == 1, "Cannot identify best epoch uniquely from saved metrics"
    best = matched.iloc[0]
    assert sha256(training / "weights" / "best.pt") == completed["weights_sha256"]
    checks = {"metric_rows_recomputed": len(metrics), "paired_intervals_recomputed": len(intervals),
              "event_match_rows_checked": len(events), "rr_scored_windows": int(registry.eligible_rr.sum()),
              "event_scored_windows": int(registry.eligible_events.sum()), "segments": len(registry),
              "prediction_rows": len(predictions), "best_epoch": int(best.epoch), "last_epoch": int(epochs.epoch.max()),
              "best_weight_metrics_match_epoch": True,
              "all_missing_signal_outputs": int(predictions.output_status.eq("zero_count_from_all_missing_signal").sum())}
    return predictions, metrics, intervals, registry, epochs, best, ready, checks


def metric_table(metrics, split):
    rows = []
    for cohort in COHORTS:
        group = metrics[metrics.cohort.eq(cohort) & metrics.split.eq(split)].set_index("arm")
        for key, label, unit, _ in METRICS:
            record = {"cohort": cohort, "队列": COHORT_LABELS[cohort], "split": split,
                      "指标": label, "单位或说明": unit, "metric_key": key}
            for arm in ARMS:
                if arm not in group.index:
                    value = None
                elif key == "exact_count_display":
                    value = f"{int(group.loc[arm, 'exact_count'])}/{int(group.loc[arm, 'scored_windows'])}"
                else:
                    value = group.loc[arm, key]
                record[arm] = value
            record["备注"] = "无完整RR参考，不评分" if group.empty else (
                "次数参考，无逐事件F1" if cohort == "short73" else
                "仅1个可评分窗，R²不可计算" if group.scored_windows.min() == 1 else "仅完整参考窗口评分")
            rows.append(record)
    return pd.DataFrame(rows)


def paired_table(intervals, split):
    rows = []
    for cohort in COHORTS:
        group = intervals[intervals.cohort.eq(cohort) & intervals.split.eq(split)].set_index("candidate")
        for key, label in (("delta_mae_bpm", "ΔMAE（候选减C0P0）"), ("ci95_low", "95%区间下限"),
                           ("ci95_high", "95%区间上限"), ("windows", "配对窗口数"), ("cow_groups", "牛号组数")):
            rows.append({"队列": COHORT_LABELS[cohort], "split": split, "指标": label,
                         "C0P0": "对照" if key == "delta_mae_bpm" else None,
                         **{arm: group.loc[arm, key] if arm in group.index else None for arm in ARMS[1:]}})
    return pd.DataFrame(rows)


def window_table(predictions, registry):
    table = registry.copy().set_index("segment_id")
    for key in ("truth_rr_bpm",):
        table[key] = predictions[predictions.arm.eq("C0P0")].set_index("segment_id")[key]
    for arm in ARMS:
        group = predictions[predictions.arm.eq(arm)].set_index("segment_id")
        for key in ("predicted_count", "predicted_rr_bpm", "absolute_rr_error_bpm", "exact_count",
                    "mean_count_accuracy", "tp", "fp", "fn", "fusion", "output_status"):
            table[f"{key}_{arm}"] = group[key]
    return table.reset_index()


def excel_value(value):
    if value is None or pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def add_sheet(workbook, name, frame, source, widths=None):
    sheet = workbook.create_sheet(name)
    sheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        sheet.append([excel_value(value) for value in row])
    sheet.freeze_panes = "D2" if len(frame.columns) > 6 else "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="244C46")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.comment = Comment(f"Source snapshot: {source}", "Codex")
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 25
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=10, color="202625")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F1F5F4")
            if isinstance(cell.value, float):
                cell.number_format = "0.000000"
    for index, column in enumerate(frame.columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(column, 22)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_title_rows = "1:1"
    return sheet


def markdown_table(frame, formats=None):
    lines = ["| " + " | ".join(map(str, frame.columns)) + " |", "|" + "---|" * len(frame.columns)]
    for row in frame.itertuples(index=False, name=None):
        cells = []
        for column, value in zip(frame.columns, row):
            if pd.isna(value):
                cell = "不适用/不可计算"
            elif isinstance(value, (float, np.floating)):
                cell = format(value, (formats or {}).get(column, ".6f"))
            else:
                cell = str(value)
            cells.append(cell.replace("|", "/").replace("\n", " "))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--replay", type=Path, default=ROOT / "replay_lindian_20261016_v1")
    parser.add_argument("--training", type=Path, default=ROOT / "training_runs" / "yolo11n_pose_lindian_20261016_v1")
    parser.add_argument("--dataset", type=Path, default=ROOT / "yolo_dataset_lindian_20261016_v1")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--include-pooled", action="store_true", help="Add clearly labelled exploratory totals and audit 73 input durations")
    args = parser.parse_args()
    replay, training, output = args.replay.resolve(), args.training.resolve(), args.output.resolve()
    if output.exists():
        raise FileExistsError(f"Preserve prior metric snapshots: {output}")
    sources = [replay / name for name in ("COMPLETED.json", "metrics.csv", "paired_intervals.csv", "predictions.csv",
               "event_matches.csv", "reference_registry.csv", "freeze.json")]
    sources += [training / "results.csv", training / "weights" / "best.pt", args.dataset / "READY.json"]
    if args.include_pooled:
        sources += [ROOT / "corpus_v1" / "segments.csv", ROOT / "corpus_v1" / "manual_review_status.csv",
                    replay / "frames.csv", replay / "frame_times.csv",
                    Path("E:/real/汇报+论文/组会/第三次工作汇报材料/cow_rr_group_meeting_20260908_v4.pptx")]
    fingerprints = [{"path": str(path), "sha256": sha256(path)} for path in sources]
    predictions, metrics, intervals, registry, epochs, best, ready, checks = check_sources(replay, training, args.dataset)
    output.mkdir(parents=True)
    yolo = pd.DataFrame([
        {"指标": label, "鼻部检测框(B)": best[f"metrics/{key}(B)"], "鼻孔关键点(P)": best[f"metrics/{key}(P)"],
         "来源": f"best.pt保存的train_metrics，与results.csv第{int(best.epoch)}轮一致"}
        for key, label in (("precision", "Precision"), ("recall", "Recall"), ("mAP50", "mAP50"), ("mAP50-95", "mAP50-95"))
    ])
    notes = [
        ("快照类型", "导入既有实验计算结果的只读快照；不含Excel公式，不修改算法、真值、窗口或默认模型。"),
        ("数据范围", f"林甸{len(registry)}片段，{ready['frames']}帧；train {ready['train']} / val {ready['val']}。R3久福暂缓。"),
        ("训练完成", f"第{int(epochs.epoch.max())}轮早停，最佳模型第{int(best.epoch)}轮；表中为最佳轮验证指标，不是最后一轮。"),
        ("YOLO指标尺度", "Precision、Recall、mAP按0至1存储；98.490%等百分比写法需乘100，不能和RR事件F1混同。"),
        ("固定输入", "同一新YOLO11n-Pose、BGR随机森林温度映射、固定半径20、置信度0.5、MAF3；不做短缺口重取。"),
        ("方法C0P0", ARM_LABELS["C0P0"]), ("方法C1P0", ARM_LABELS["C1P0"]),
        ("方法C0P1", ARM_LABELS["C0P1"]), ("方法C1P1", ARM_LABELS["C1P1"]),
        ("质量选择含义", "整曲线候选可含左、右及融合曲线，不能简化成只选左右单鼻孔。"),
        ("RR定义", "RR=60N/T；73短片用各自实际6至20秒时长；49窗用冻结播放时间约30秒；P1/P2/P3用冻结30秒。"),
        ("R²定义", "1-SSE/SST，不是Pearson相关系数平方；单窗或参考无方差时留空，不写0。"),
        ("平均计数准确度", "逐窗max(0,1-|预测次数-参考次数|/参考次数)的平均；参考0次时用预测是否也为0。不是事件匹配准确率。"),
        ("事件F1", "2TP/(2TP+FP+FN)，按±0.30秒一对一匹配汇总；B73-I没有逐事件参考，留空。"),
        ("RR参考", "可评分73/73、47/49、P1 5/8、P2 6/16、P3 24/53；不完整参考及两段事后36次排除。R1分歧未替换P2。"),
        ("覆盖说明", "纳入/可评分窗数是参考覆盖；不是拒测覆盖。固定规则即使信号全缺失仍输出0，不能解释为生理0次或可靠测量。"),
        ("分组说明", "all含train和val。train/val按YOLO牛号组划分；这些RR窗口已用于开发，val也不是未见RR方法确认或跨牧场外测。"),
        ("配对区间", "ΔMAE=候选减C0P0，负数为改善；按牛号组bootstrap 10000次，seed20260929；单牛号组区间不代表群体推断。"),
        ("未采用结论", "47窗合并候选MAE区间跨0且完全计数下降，P3退化；本表不改变默认方法，也不据此宣称创新独立证实。"),
        ("以后结果格式", "每次实验另存版本目录；指标作行、方法作列；每队列分别列all/train/val，并保留逐窗预测、参考和输入哈希。"),
        ("空值说明", "原始表空值保持为空；P2验证牛号组无完整参考，因此不评分；P1和P3验证组各仅1窗，R²不可计算。"),
        ("权重路径", str(training / "weights" / "best.pt")),
        ("权重SHA256", next(item["sha256"] for item in fingerprints if item["path"] == str(training / "weights" / "best.pt"))),
        ("结果目录", str(replay)),
    ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "说明", pd.DataFrame(notes, columns=["项目", "说明"]), replay / "freeze.json", {"项目": 25, "说明": 115})
    for row in workbook["说明"].iter_rows(min_row=2):
        workbook["说明"].row_dimensions[row[0].row].height = 42
    add_sheet(workbook, "YOLO最佳轮验证", yolo, training / "results.csv", {"指标": 23, "来源": 76})
    report = ["# 重训后指标横向对照\n", "## 口径\n", *[f"- **{key}**：{value}" for key, value in notes],
              "\n## YOLO最佳轮验证\n", markdown_table(yolo.drop(columns="来源")),
              "\n此表数值为0至1；下面RR表的平均计数准确度以百分比显示，其他R²/F1仍为0至1。\n"]
    for split in SPLITS:
        wide = metric_table(metrics, split)
        wide.to_csv(output / f"rr_metrics_{split}_wide.csv", index=False, encoding="utf-8-sig")
        sheet = add_sheet(workbook, f"RR_{split}", wide.drop(columns=["cohort", "metric_key"]), replay / "metrics.csv",
                          {"队列": 23, "指标": 28, "单位或说明": 39, "备注": 42, **{arm: 22 for arm in ARMS}})
        for index, key in enumerate(wide.metric_key, 2):
            number_format = next(fmt for name, _, _, fmt in METRICS if name == key)
            for arm in ARMS:
                sheet.cell(index, list(wide.drop(columns=["cohort", "metric_key"]).columns).index(arm) + 1).number_format = number_format
        paired = paired_table(intervals, split)
        add_sheet(workbook, f"配对区间_{split}", paired, replay / "paired_intervals.csv", {"队列": 23, "指标": 35})
        report.append(f"\n## {SPLIT_LABELS[split]}\n")
        for cohort in COHORTS:
            group = wide[wide.cohort.eq(cohort)].astype(object).copy()
            accuracy = group.metric_key.eq("mean_count_accuracy")
            for arm in ARMS:
                group.loc[accuracy, arm] = group.loc[accuracy, arm].map(lambda value: f"{value:.2%}" if pd.notna(value) else None)
            report += [f"\n### {COHORT_LABELS[cohort]}\n", markdown_table(group[["指标", *ARMS]])]
        report += ["\n### 相对C0P0的MAE配对差及区间\n", markdown_table(paired)]
    windows = window_table(predictions, registry)
    windows.to_csv(output / "window_predictions_wide.csv", index=False, encoding="utf-8-sig")
    yolo.to_csv(output / "yolo_best_epoch_metrics.csv", index=False, encoding="utf-8-sig")
    add_sheet(workbook, "逐窗四组明细", windows, replay / "predictions.csv", {"segment_id": 52, "cow_group": 27})
    add_sheet(workbook, "原始RR汇总", metrics, replay / "metrics.csv")
    add_sheet(workbook, "原始配对区间", intervals, replay / "paired_intervals.csv")
    add_sheet(workbook, "逐轮训练记录", epochs, training / "results.csv")
    add_sheet(workbook, "来源哈希", pd.DataFrame(fingerprints), "SHA256 of imported files", {"path": 120, "sha256": 73})
    if args.include_pooled:
        from pooled_rr_metrics import augment

        checks.update(augment(workbook, report, predictions, metrics, registry, replay, output, ROOT,
                              add_sheet, markdown_table, sha256))
    workbook.save(output / "metric_comparison.xlsx")
    (output / "指标对照说明.md").write_text("\n".join(report) + "\n", encoding="utf-8")
    reopened = load_workbook(output / "metric_comparison.xlsx", data_only=False)
    assert reopened["逐窗四组明细"].max_row == len(registry) + 1
    assert reopened["原始RR汇总"].max_row == len(metrics) + 1
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}
    for sheet in reopened:
        for row in sheet:
            for cell in row:
                assert cell.data_type not in ("f", "e"), f"Unexpected formula or error: {sheet.title}!{cell.coordinate}"
                assert not isinstance(cell.value, str) or cell.value not in errors
                assert not isinstance(cell.value, float) or math.isfinite(cell.value)
    for split in SPLITS:
        frame = metric_table(metrics, split)
        sheet = reopened[f"RR_{split}"]
        exported_columns = list(frame.drop(columns=["cohort", "metric_key"]).columns)
        for row_number, row in enumerate(frame.to_dict("records"), 2):
            for arm in ARMS:
                actual = sheet.cell(row_number, exported_columns.index(arm) + 1).value
                expected = excel_value(row[arm])
                if isinstance(expected, (int, float)):
                    same(actual, expected, "Workbook metric")
                else:
                    assert actual == expected
    for item in fingerprints:
        assert sha256(Path(item["path"])) == item["sha256"], "Source changed while exporting"
    checks.update(status="passed", workbook_sheets=reopened.sheetnames,
                  workbook_formula_count=0, workbook_error_count=0,
                  source_files_unchanged=True, generated_at=datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
                  sources=fingerprints)
    (output / "verification.json").write_text(json.dumps(checks, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in checks.items() if key != "sources"}, ensure_ascii=False, indent=2))
    print(f"OUTPUT {output}")


if __name__ == "__main__":
    main()
